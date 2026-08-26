"""
tests/test_e2e_integration.py

Phase 15: single end-to-end test running the exact assignment demo
scenario start to finish: upload DOCX+PPTX templates -> /chat triggers
analysis+research+RAG+generation+validation -> download artifacts ->
conversational edits -> version history -> trace lookup.
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import hashlib
import shutil
import numpy as np
import pytest
from docx import Document as DocxDocument
from pptx import Presentation
from fastapi.testclient import TestClient

from backend.main import app
from backend.config import settings
from backend.services.citation_service import citation_registry
from backend.services.embedding_service import EmbeddingService, EMBEDDING_DIM
from backend.services.vector_store import reset_vector_store_cache

client = TestClient(app)
FIXTURES = Path(__file__).resolve().parent / "fixtures"


def _fake_embed_texts(self, texts):
    """Same deterministic offline embedder used in tests/test_rag.py — avoids
    depending on Hugging Face network access for the model download."""
    vectors = []
    for text in texts:
        vec = np.zeros(EMBEDDING_DIM, dtype=np.float32)
        for word in text.lower().split():
            idx = int(hashlib.md5(word.encode()).hexdigest(), 16) % EMBEDDING_DIM
            vec[idx] += 1.0
        norm = np.linalg.norm(vec)
        if norm > 0:
            vec = vec / norm
        vectors.append(vec.tolist())
    return vectors


@pytest.fixture(autouse=True)
def _reset(monkeypatch):
    monkeypatch.setattr(EmbeddingService, "embed_texts", _fake_embed_texts)
    reset_vector_store_cache()
    faiss_dir = settings.resolve_path(settings.knowledge_base_dir) / "faiss_index"
    if faiss_dir.exists():
        shutil.rmtree(faiss_dir)
    citation_registry.reset()
    yield
    reset_vector_store_cache()


def test_full_assignment_demo_scenario_end_to_end():
    # 1. Upload DOCX + PPTX templates
    with open(FIXTURES / "sample.docx", "rb") as f:
        up1 = client.post("/upload", files={"file": ("Company_Proposal.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    with open(FIXTURES / "sample.pptx", "rb") as f:
        up2 = client.post("/upload", files={"file": ("Company_Template.pptx", f, "application/vnd.openxmlformats-officedocument.presentationml.presentation")})
    assert up1.status_code == 200 and up2.status_code == 200
    docx_id, pptx_id = up1.json()["file"]["file_id"], up2.json()["file"]["file_id"]

    # 2-11. Full orchestration: analyze, research, RAG, generate, validate, version, trace
    chat_resp = client.post("/chat", json={
        "message": "Research the latest Generative AI trends and create a proposal and 12-slide presentation using the same tone and style as the uploaded files.",
        "file_ids": [docx_id, pptx_id],
    })
    assert chat_resp.status_code == 200
    result = chat_resp.json()

    assert result["document_analysis"] is not None
    assert result["ppt_analysis"] is not None
    assert result["research"] is not None
    assert result["rag_results"] is not None
    assert result["generated_docx_artifact_id"] is not None
    assert result["generated_pptx_artifact_id"] is not None
    assert result["docx_validation_status"] == "PASS"
    assert result["pptx_validation_status"] == "PASS"
    assert result["trace_id"] is not None

    docx_artifact_id = result["generated_docx_artifact_id"]
    pptx_artifact_id = result["generated_pptx_artifact_id"]

    # 12. Download both real artifacts and verify they actually open
    docx_dl = client.get(f"/artifact/{docx_artifact_id}/download")
    assert docx_dl.status_code == 200
    pptx_dl = client.get(f"/artifact/{pptx_artifact_id}/download")
    assert pptx_dl.status_code == 200

    from backend.services import artifact_store
    docx_meta = artifact_store.load_meta(docx_artifact_id)
    pptx_meta = artifact_store.load_meta(pptx_artifact_id)
    doc = DocxDocument(docx_meta["versions"][0]["file_path"])
    prs = Presentation(pptx_meta["versions"][0]["file_path"])
    assert len(doc.paragraphs) > 0
    assert len(prs.slides) == 12  # requested slide count honored

    # 13. Traceability panel: full trace retrievable
    trace_resp = client.get(f"/trace/{result['trace_id']}")
    assert trace_resp.status_code == 200
    trace = trace_resp.json()
    agent_names = {c["agent"] for c in trace["agent_calls"]}
    assert {"document_analyzer", "ppt_analyzer", "web_research", "rag_agent", "document_generator", "ppt_generator"} <= agent_names

    # 14. Conversational editing: "Add an executive summary" (test the exact demo follow-up)
    edit1 = client.post(f"/artifact/{docx_artifact_id}/edit", json={"request": "Add an executive summary"})
    assert edit1.status_code == 200
    assert edit1.json()["new_version"] == 2

    # 15. "Make the presentation more concise"
    edit2 = client.post(f"/artifact/{pptx_artifact_id}/edit", json={"request": "Make the presentation more concise"})
    assert edit2.status_code == 200
    assert edit2.json()["new_version"] == 2

    # 16. "Add a competitive analysis section"
    edit3 = client.post(f"/artifact/{docx_artifact_id}/edit", json={"request": "Add a competitive analysis section"})
    assert edit3.status_code == 200
    assert edit3.json()["new_version"] == 3

    # 17. Version history reflects all edits, each version preserved
    versions_resp = client.get(f"/artifact/{docx_artifact_id}/versions")
    versions = versions_resp.json()
    assert len(versions) == 3
    for v in versions:
        assert Path(v["file_path"]).exists()

    # 18. Doc -> PPT conversion works on the edited artifact (no slide_count
    # forced here — the edited doc now has more sections than a small fixed
    # count could hold, and the generator correctly refuses to silently drop
    # content to hit an under-sized target; that's tested explicitly in
    # test_ppt_generation.py's target_slide_count tests).
    conv_resp = client.post("/convert/document-to-presentation", json={"artifact_id": docx_artifact_id})
    assert conv_resp.status_code == 200
    assert conv_resp.json()["validation_status"] == "PASS"
